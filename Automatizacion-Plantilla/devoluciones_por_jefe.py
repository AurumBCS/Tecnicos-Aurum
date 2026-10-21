"""Control diario de devoluciones al almacen por jefe de equipo.

Entrada: los 3 Excel de la app OV -> Inventario Transacciones -> Devoluciones al almacen
(Borrador, Completado, Recibo Pendiente) y el cuadrante (CSV del script de la app).
Lo usa ov_devoluciones.py (Lambda). Uso local:
  python devoluciones_por_jefe.py <borrador> <completado> <pendiente> <cuadrante.csv> <carpeta> [AAAA-MM-DD]

Reglas (acordadas con Gustavo, 2026-10-17):
  - Objetivo: 2 devoluciones al mes por tecnico. Cuenta como hecha una devolucion
    Completada o en Recibo Pendiente (ya enviada); la fecha es "Creado".
  - Mas de 7 dias sin devolver = aviso; mas de 15 = grave, hacerla ya. Si nunca ha
    devuelto se cuenta desde su Inicio TOA.
  - Borrador del mismo dia = la hizo en fisico pero no la completo en digital.
    Borrador de dias anteriores = completarla, o borrarlo y crear otro cuando toque.
  - Recibo Pendiente: se listan con los dias que llevan esperando a que el almacen
    las reciba (luego pasan a Completado).
  - Tecnicos: mismos que el control de recuentos (5 TEC del cuadrante, fila > 24, sin
    baja y con Inicio TOA); Nicola Morsa (302154) -> Gustavo Perez (GD5381).
"""
import base64
import html as _html
import json
import os
import sys
from datetime import date, datetime

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

AVISO_DIAS, GRAVE_DIAS, OBJETIVO_MES = 7, 15, 2
REASIGNAR_JDE = {"302154": "GD5381"}
# Jefes a los que no se saluda por el nombre (GD5381 = Gustavo, que recibe su propio correo con Nicola en CCO)
SALUDO_NEUTRO = {"GD5381"}


def saludo(jde, nombre):
    return "Buenos días," if jde in SALUDO_NEUTRO else f"Hola {nombre.split()[0]},"
GRAVE = f"Más de {GRAVE_DIAS} días sin devolver"
AVISO = f"Más de {AVISO_DIAS} días sin devolver"
OK = "Al día"


def _matricula(subinv):
    """'357633DF DF Krystiam ...' -> '357633' ; '358719 Ramón ...' -> '358719'."""
    tok = str(subinv or "").strip().split(" ")[0].upper()
    return tok[:-2] if tok.endswith("DF") and len(tok) > 2 else tok


def _fecha(v):
    if isinstance(v, (datetime, pd.Timestamp)):
        return v.date()
    if isinstance(v, date):
        return v
    d = pd.to_datetime(str(v).strip(), dayfirst=True, errors="coerce")
    return None if pd.isna(d) else d.date()


def leer_devoluciones(ruta):
    """Excel descargado de Ver Devoluciones (o CSV ';' de prueba) -> Devolución, Tipo, Creado, Matrícula."""
    if str(ruta).lower().endswith(".csv"):
        df = pd.read_csv(ruta, sep=";", dtype=str)
    else:
        crudo = pd.read_excel(ruta, header=None, dtype=object)
        fila = next(i for i in range(min(6, len(crudo)))
                    if "Devolución" in [str(x).strip() for x in crudo.iloc[i]])
        df = crudo.iloc[fila + 1:].copy()
        df.columns = [str(x).strip() for x in crudo.iloc[fila]]
    df = df.dropna(subset=["Devolución"])
    return pd.DataFrame({
        "Devolución": df["Devolución"].astype(str).str.strip(),
        "Tipo": df["Tipo Devolución"].astype(str).str.strip(),
        "Creado": df["Creado"].map(_fecha),
        "Matrícula": df["Subinventario Origen"].map(_matricula),
    })


def leer_cuadrante(ruta, exigir_toa=True):
    """CSV del endpoint ?type=recuentos -> (tecnicos activos, nombres por matricula)."""
    c = pd.read_csv(ruta, dtype=str).apply(lambda s: s.str.strip()).replace("", pd.NA)
    c["Fila"] = pd.to_numeric(c["Fila"]) if "Fila" in c else range(3, 3 + len(c))
    if "Columna O" in c:
        c["Inicio TOA"] = c["Columna O"]
    toa = pd.to_datetime(c["Inicio TOA"], format="%d/%m/%Y", errors="coerce")
    toa = toa.fillna(pd.to_datetime(pd.to_numeric(c["Inicio TOA"], errors="coerce"), unit="D", origin="1899-12-30"))
    c["TOA"] = toa.dt.date
    c = c.dropna(subset=["Nombre(S)"])
    c["Matricula"] = c["Matricula"].astype(str).str.strip().str.upper()
    c["JDE"] = c["JDE"].astype(str).str.strip().str.upper().str.replace(r"\.0$", "", regex=True).replace(REASIGNAR_JDE)
    c["Nombre"] = (c["Nombre(S)"].str.strip() + " " + c["Apellidos(S)"].fillna("").str.strip()).str.strip()
    nombres = c.set_index("Matricula")["Nombre"].to_dict()
    activos = c[(c["Fila"] > 24) & (c["Jerarquia"] == "5 TEC") & c["F Baja"].isna()]
    if exigir_toa:
        activos = activos[activos["TOA"].notna()]
    return activos, nombres


def calcular(borr, comp, pend, activos, hoy):
    hechas = pd.concat([comp, pend])
    filas = []
    for _, t in activos.iterrows():
        h = hechas[hechas["Matrícula"] == t.Matricula]
        ultima = max(h["Creado"].dropna(), default=None)
        mes = int(sum(1 for d in h["Creado"].dropna() if (d.year, d.month) == (hoy.year, hoy.month)))
        dias = (hoy - (ultima or t.TOA)).days
        estado = GRAVE if dias > GRAVE_DIAS else AVISO if dias > AVISO_DIAS else OK
        b = borr[borr["Matrícula"] == t.Matricula]
        filas.append({
            "JDE": t.JDE, "Matrícula": t.Matricula, "Técnico": t.Nombre, "Provincia": t.Provincia,
            "Inicio TOA": t.TOA, "Última devolución": ultima,
            "Días sin devolver": dias, f"Devoluciones este mes (objetivo {OBJETIVO_MES})": mes,
            "Borradores sin enviar": len(b), "Recibo pendiente": int((pend["Matrícula"] == t.Matricula).sum()),
            "Estado": estado,
        })
    df = pd.DataFrame(filas)
    orden = {GRAVE: 0, AVISO: 1, OK: 2}
    df = df.assign(_o=df.Estado.map(orden)).sort_values(["JDE", "_o", "Días sin devolver"], ascending=[True, True, False]).drop(columns="_o")

    jde = activos.set_index("Matricula")["JDE"].to_dict()
    nombre = activos.set_index("Matricula")["Nombre"].to_dict()

    def lista(d, situacion):
        d = d[d["Matrícula"].isin(jde)].copy()
        d.insert(0, "JDE", d["Matrícula"].map(jde))
        d.insert(2, "Técnico", d["Matrícula"].map(nombre))
        d["Días"] = d["Creado"].map(lambda f: (hoy - f).days if f else None)
        d["Situación"] = d.apply(situacion, axis=1)
        return d.sort_values(["JDE", "Días"], ascending=[True, False])

    borradores = lista(borr, lambda r: "Creado hoy: falta completarla en digital" if r["Días"] == 0
                       else f"Hace {r['Días']} día{'s' if r['Días'] != 1 else ''}: completarla, o borrar el borrador "
                            "y crear otro cuando toque")
    pendientes = lista(pend, lambda r: f"Esperando recepción en almacén desde hace {r['Días']} días")
    fuera = pd.concat([borr.assign(Estado="Borrador"), pend.assign(Estado="Recibo pendiente")])
    fuera = fuera[~fuera["Matrícula"].isin(jde)]
    return df, borradores, pendientes, fuera


# --- Excel ---
HDR = PatternFill("solid", fgColor="1F3864")
COLORES = {GRAVE: "F8CBAD", AVISO: "FFE699", OK: "C6EFCE"}


def escribir(ws, data, titulo, hoy, nota=None):
    ws.cell(1, 1, titulo).font = Font(bold=True, size=13)
    nota = nota or (f"Objetivo: {OBJETIVO_MES} devoluciones al mes. "
                    f"Más de {AVISO_DIAS} días sin devolver = aviso; más de {GRAVE_DIAS} = hacerla ya.")
    ws.cell(2, 1, f"Fecha del informe: {hoy:%d/%m/%Y} · {nota}").font = Font(italic=True, size=9)
    cols = list(data.columns)
    for j, col in enumerate(cols, 1):
        cell = ws.cell(4, j, col)
        cell.font, cell.fill = Font(bold=True, color="FFFFFF"), HDR
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for i, row in enumerate(data.itertuples(index=False), 5):
        for j, v in enumerate(row, 1):
            if v is None or (isinstance(v, float) and pd.isna(v)):
                v = None
            cell = ws.cell(i, j, v)
            if isinstance(v, (date, datetime)):
                cell.number_format = "DD/MM/YYYY"
        if "Estado" in cols and row[cols.index("Estado")] in COLORES:
            ws.cell(i, cols.index("Estado") + 1).fill = PatternFill("solid", fgColor=COLORES[row[cols.index("Estado")]])
    for j, col in enumerate(cols, 1):
        largo = max([len(str(col))] + [len(str(v)) for v in data.iloc[:, j - 1].head(200)])
        ws.column_dimensions[get_column_letter(j)].width = min(max(10, largo + 2), 60)
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A4:{get_column_letter(len(cols))}{4 + len(data)}"


# --- Correo HTML ---
def _fmt(v):
    if isinstance(v, (date, datetime)):
        return v.strftime("%d/%m/%Y")
    return "" if v is None or (isinstance(v, float) and pd.isna(v)) else str(v)


def _tabla(g, titulo, fondo, borde, cols):
    e = _html.escape
    th = 'style="text-align:left;padding:6px 8px;font-size:12px;color:#555;border-bottom:1px solid #ddd"'
    td = 'style="padding:6px 8px;font-size:13px;border-bottom:1px solid #eee;vertical-align:top"'
    cab = f"<th {th}>Técnico</th>" + "".join(f"<th {th}>{e(c)}</th>" for c, _ in cols)
    filas = "".join(
        f"<tr><td {td}><b>{e(str(t['Técnico']))}</b><br><span style=\"color:#888;font-size:11px\">{e(t['Matrícula'])}</span></td>"
        + "".join(f"<td {td}>{e(f(t))}</td>" for _, f in cols) + "</tr>"
        for _, t in g.iterrows())
    return (f'<table role="presentation" width="100%" cellpadding="0" cellspacing="0" '
            f'style="border-collapse:collapse;margin:18px 0 4px;background:{fondo};border-left:4px solid {borde}">'
            f'<tr><td style="padding:8px 10px;font-size:14px;font-weight:bold;color:{borde}">{e(titulo)} ({len(g)})</td></tr></table>'
            f'<table role="presentation" width="100%" cellpadding="0" cellspacing="0" style="border-collapse:collapse">'
            f"<tr>{cab}</tr>{filas}</table>")


MES = f"Devoluciones este mes (objetivo {OBJETIVO_MES})"


def _secciones(g, b, p):
    sin = [("Última devolución", lambda t: _fmt(t["Última devolución"]) or f"Ninguna (Inicio TOA {_fmt(t['Inicio TOA'])})"),
           ("Días sin devolver", lambda t: str(t["Días sin devolver"])),
           ("Este mes", lambda t: f"{t[MES]} de {OBJETIVO_MES}")]
    return [  # (título, corto, fondo, borde, filas, columnas)
        ("Borradores sin enviar", "Borradores", "#FFF1E6", "#D35400", b,
         [("Devolución", lambda t: t["Devolución"]), ("Creado", lambda t: _fmt(t["Creado"])),
          ("Qué hacer", lambda t: t["Situación"])]),
        (f"{GRAVE} – hacer la devolución ya", f"+{GRAVE_DIAS} días", "#FDECEA", "#C0392B", g[g.Estado == GRAVE], sin),
        (AVISO, f"+{AVISO_DIAS} días", "#FFF8E1", "#E0A800", g[g.Estado == AVISO], sin),
        ("Recibo pendiente en almacén", "Recibo pendiente", "#EAF2FB", "#2E75B6", p,
         [("Devolución", lambda t: t["Devolución"]), ("Enviada", lambda t: _fmt(t["Creado"])),
          ("Días esperando", lambda t: str(t["Días"]))]),
        (OK, OK, "#EAF6EC", "#2E8B57", g[g.Estado == OK],
         [("Última devolución", lambda t: _fmt(t["Última devolución"])), ("Este mes", lambda t: f"{t[MES]} de {OBJETIVO_MES}")]),
    ]


def generar(rutas, cuadrante, carpeta, hoy):
    """rutas = {"D": borrador, "C": completado, "P": pendiente}. Devuelve (Excel general, carpeta correos, resumen)."""
    borr, comp, pend = (leer_devoluciones(rutas[k]) for k in "DCP")
    activos, nombres = leer_cuadrante(cuadrante)
    df, borradores, pendientes, fuera = calcular(borr, comp, pend, activos, hoy)

    os.makedirs(carpeta, exist_ok=True)
    general = os.path.join(carpeta, f"Devoluciones_por_Jefe_{hoy:%Y-%m-%d}.xlsx")
    res = df.groupby("JDE").agg(Técnicos=("Matrícula", "count")).reset_index()
    for est in (GRAVE, AVISO, OK):
        res[est] = res.JDE.map(df[df.Estado == est].groupby("JDE").size()).fillna(0).astype(int)
    res["Borradores sin enviar"] = res.JDE.map(borradores.groupby("JDE").size()).fillna(0).astype(int)
    res["Recibo pendiente"] = res.JDE.map(pendientes.groupby("JDE").size()).fillna(0).astype(int)
    res.insert(1, "Jefe de equipo", res.JDE.map(lambda j: nombres.get(j, j)))
    wb = Workbook()
    wb.active.title = "Resumen"
    escribir(wb.active, res, "Control de devoluciones por jefe de equipo", hoy)
    escribir(wb.create_sheet("Técnicos"), df, "Técnicos: días sin devolver", hoy)
    escribir(wb.create_sheet("Borradores"), borradores, "Borradores sin enviar", hoy)
    escribir(wb.create_sheet("Recibo pendiente"), pendientes, "Recibo pendiente en almacén", hoy)
    escribir(wb.create_sheet("Fuera del control"), fuera,
             "Borradores / recibos de personas fuera del control (JDE, coordinadores, bajas o no activas)", hoy)
    wb.save(general)

    dir_correos = os.path.join(carpeta, f"correos_{hoy:%Y-%m-%d}")
    os.makedirs(dir_correos, exist_ok=True)
    manifiesto = []
    for jde, g in df.groupby("JDE"):
        nombre = nombres.get(jde, jde)
        b, p = borradores[borradores.JDE == jde], pendientes[pendientes.JDE == jde]
        archivo = f"Devoluciones_{nombre.replace(' ', '_')}_{hoy:%Y-%m-%d}.xlsx"
        wbj = Workbook()
        wbj.active.title = "Mi equipo"
        escribir(wbj.active, g.drop(columns="JDE"), f"Jefe de equipo: {nombre} ({jde})", hoy)
        escribir(wbj.create_sheet("Borradores"), b.drop(columns="JDE"), "Borradores sin enviar", hoy)
        escribir(wbj.create_sheet("Recibo pendiente"), p.drop(columns="JDE"), "Recibo pendiente en almacén", hoy)
        ruta = os.path.join(dir_correos, archivo)
        wbj.save(ruta)

        bloques, cajas = [], []
        for titulo, corto, fondo, borde, sub, cols in _secciones(g, b, p):
            cajas.append(f'<td align="center" style="padding:8px 4px;background:{fondo};border-top:3px solid {borde}">'
                         f'<div style="font-size:22px;font-weight:bold;color:{borde}">{len(sub)}</div>'
                         f'<div style="font-size:11px;color:#555">{_html.escape(corto)}</div></td>')
            if len(sub):
                bloques.append(_tabla(sub, titulo, fondo, borde, cols))
        cuerpo = f"""<div style="font-family:Arial,Helvetica,sans-serif;color:#222;max-width:680px">
<p style="font-size:14px">{_html.escape(saludo(jde, nombre))}</p>
<p style="font-size:14px">Te paso el control de devoluciones al almacén de tu equipo a fecha <b>{hoy:%d/%m/%Y}</b>.
Cada técnico debe hacer <b>{OBJETIVO_MES} devoluciones al mes</b>. Con más de {AVISO_DIAS} días sin devolver hay que
organizarla, y con más de {GRAVE_DIAS} días hay que hacerla ya. Los borradores sin enviar hay que completarlos en digital,
o borrarlos y crear otro cuando toque.</p>
<table role="presentation" width="100%" cellpadding="0" cellspacing="4" style="margin:10px 0"><tr>{''.join(cajas)}</tr></table>
{''.join(bloques)}
<p style="font-size:14px;margin-top:20px">Te adjunto el Excel con el detalle completo. Avísame si algún dato no cuadra.</p>
<p style="font-size:14px">Un saludo,</p>
<p style="font-size:14px;margin-top:4px">Gustavo Perez<br>Aurum BC Security SL</p></div>"""
        graves = int((g.Estado == GRAVE).sum())
        manifiesto.append({
            "jde": jde, "nombre": nombre, "asunto": f"Control de devoluciones de tu equipo - {hoy:%d/%m/%Y}",
            "html": cuerpo, "archivo": archivo,
            "texto": f"Control de devoluciones de tu equipo a {hoy:%d/%m/%Y}. Borradores sin enviar: {len(b)}. "
                     f"Más de {GRAVE_DIAS} días sin devolver: {graves}. Detalle en el Excel adjunto.",
            "xlsx_b64": base64.b64encode(open(ruta, "rb").read()).decode(),
        })
        with open(ruta.replace(".xlsx", ".html"), "w", encoding="utf-8") as fh:
            fh.write('<!doctype html><html lang="es"><head><meta charset="utf-8">'
                     '<meta name="viewport" content="width=device-width,initial-scale=1">'
                     f'<title>Correo {_html.escape(nombre)}</title></head>'
                     f'<body style="margin:0;padding:16px;background:#fff">{cuerpo}</body></html>')
    with open(os.path.join(dir_correos, "correos.json"), "w", encoding="utf-8") as fh:
        json.dump(manifiesto, fh, ensure_ascii=False)

    resumen = (df.Estado.value_counts().to_string() + f"\nBorradores sin enviar: {len(borradores)}"
               f"\nRecibo pendiente: {len(pendientes)}\nFuera del control: {len(fuera)}\n\n"
               + res.drop(columns="JDE").to_string(index=False))
    return general, dir_correos, resumen


if __name__ == "__main__":
    d, c, p, cuadrante, carpeta = sys.argv[1:6]
    hoy = date.fromisoformat(sys.argv[6]) if len(sys.argv) > 6 else date.today()
    general, correos, resumen = generar({"D": d, "C": c, "P": p}, cuadrante, carpeta, hoy)
    print(resumen)
    print("correos preparados en:", correos)
