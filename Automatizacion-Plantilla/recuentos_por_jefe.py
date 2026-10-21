"""Control de recuentos (Auditoría + Periódico) por jefe de equipo.

Uso: python recuentos_por_jefe.py <Tabla Ver Conteos.xlsx> <cuadrante.xlsx> <salida.xlsx>
Regla: 2 recuentos al año por técnico. Si no tiene ninguno y lleva más de 2 meses
desde "Inicio TOA", queda PENDIENTE para su jefe de equipo.
"""
import sys
from datetime import datetime

import pandas as pd
from dateutil.relativedelta import relativedelta
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

conteos_path, cuadrante_path, salida = sys.argv[1:4]
HOY = (datetime.fromisoformat(sys.argv[4]) if len(sys.argv) >= 5
       else datetime.combine(datetime.today().date(), datetime.min.time()))
LIMITE_TOA = HOY - relativedelta(months=2)
TIPOS = ["Auditoría", "Periódico"]

# --- Conteos ---
k = pd.read_excel(conteos_path)
k = k[k["Tipo de Conteo"].isin(TIPOS)].copy()
k["Matricula"] = k["Técnico"].str.split(" ").str[0]
k = k[k["Matricula"] != "GE"]  # conteos genéricos de la organización, sin técnico

# --- Cuadrante ---
if cuadrante_path.lower().endswith(".csv"):
    # CSV del endpoint ?type=recuentos del script de la app (fechas d/m/aaaa)
    c = pd.read_csv(cuadrante_path, dtype=str).apply(lambda s: s.str.strip())
    c = c.replace("", pd.NA)
    # Fila real del Sheets (encabezados en la fila 2); si el endpoint aún no la manda, se deduce por orden
    c["Fila"] = pd.to_numeric(c["Fila"]) if "Fila" in c else range(3, 3 + len(c))
    # Inicio TOA: siempre la columna O del cuadrante
    if "Columna O" in c:
        c["Inicio TOA"] = c["Columna O"]
    toa = pd.to_datetime(c["Inicio TOA"], format="%d/%m/%Y", errors="coerce")
    serie = pd.to_numeric(c["Inicio TOA"], errors="coerce")  # celdas guardadas como número de serie de Excel
    toa = toa.fillna(pd.to_datetime(serie, unit="D", origin="1899-12-30"))
    c["Inicio TOA"] = [d.to_pydatetime() if pd.notna(d) else v for d, v in zip(toa, c["Inicio TOA"])]
else:
    c = pd.read_excel(cuadrante_path, dtype={"Matricula": str, "JDE": str})
    c["Fila"] = range(2, 2 + len(c))
c = c.dropna(subset=["Nombre(S)"])
c["Matricula"] = c["Matricula"].astype(str).str.strip().str.upper()
c["JDE"] = c["JDE"].astype(str).str.strip().str.upper().str.replace(r"\.0$", "", regex=True)
# Equipos reasignados: los técnicos de Nicola Morsa (302154) los lleva Gustavo Pérez (GD5381)
REASIGNAR_JDE = {"302154": "GD5381"}
# Jefes a los que no se saluda por el nombre en el correo (GD5381 = Gustavo, con Nicola en CCO)
SALUDO_NEUTRO = {"GD5381"}
c["JDE"] = c["JDE"].replace(REASIGNAR_JDE)
c["Nombre"] = (c["Nombre(S)"].str.strip() + " " + c["Apellidos(S)"].fillna("").str.strip()).str.strip()
nombres = c.set_index("Matricula")["Nombre"].to_dict()
# Las filas 1-24 del cuadrante (encabezados + PM, ADM, COORD y JDE) no entran en el control
activos = c[(c["Fila"] > 24) & (c["Jerarquia"] == "5 TEC") & c["F Baja"].isna()].copy()
# Sin fecha de Inicio TOA (columna O) = nuevo ingreso: no entran en el control ni en los correos
_tiene_toa = activos["Inicio TOA"].map(lambda v: isinstance(v, datetime))
nuevo_ingreso = activos[~_tiene_toa].copy()
activos = activos[_tiene_toa].copy()


def fecha_toa(v):
    return v if isinstance(v, datetime) else None


def fechas(mat, tipo):
    f = k[(k.Matricula == mat) & (k["Tipo de Conteo"] == tipo)]["Creado"].sort_values()
    return list(f)


filas = []
for _, t in activos.iterrows():
    aud, per = fechas(t.Matricula, "Auditoría"), fechas(t.Matricula, "Periódico")
    todas = sorted(aud + per)
    toa = fecha_toa(t["Inicio TOA"])
    ultimo = todas[-1] if todas else None
    if len(todas) >= 2:
        estado = "OK - 2 o más recuentos"
    elif len(todas) == 1:
        estado = "Falta 2º recuento - VENCIDO" if todas[0] + relativedelta(months=6) < HOY else "Falta 2º recuento"
    elif toa is None:
        estado = "Sin recuento - sin fecha Inicio TOA"
    elif toa <= LIMITE_TOA:
        estado = "PENDIENTE - sin recuento y +2 meses en TOA"
    else:
        estado = "Aún no toca (menos de 2 meses en TOA)"
    if ultimo is not None:
        proximo = ultimo + relativedelta(months=6)
    elif toa is not None:
        proximo = max(toa + relativedelta(months=2), HOY) if toa <= LIMITE_TOA else toa + relativedelta(months=2)
    else:
        proximo = None
    filas.append({
        "JDE": t.JDE, "Matrícula": t.Matricula, "Técnico": t.Nombre, "Provincia": t.Provincia,
        "Inicio TOA": toa or (t["Inicio TOA"] if isinstance(t["Inicio TOA"], str) else "Sin dato"),
        "Nº Auditorías": len(aud), "Fechas Auditoría": ", ".join(d.strftime("%d/%m/%Y") for d in aud),
        "Nº Periódicos": len(per), "Fechas Periódico": ", ".join(d.strftime("%d/%m/%Y") for d in per),
        "Total recuentos": len(todas), "Último recuento": ultimo,
        "Próximo recuento antes de": proximo, "Estado": estado,
    })
df = pd.DataFrame(filas)
ORDEN = {"PENDIENTE - sin recuento y +2 meses en TOA": 0, "Falta 2º recuento - VENCIDO": 0.5, "Falta 2º recuento": 1,
         "Aún no toca (menos de 2 meses en TOA)": 3,
         "OK - 2 o más recuentos": 4}
df["_o"] = df.Estado.map(ORDEN)
df = df.sort_values(["JDE", "_o", "Técnico"]).drop(columns="_o")

# Conteos de técnicos que no están en el cuadrante (altas recientes, etc.)
fuera = k[~k.Matricula.isin(activos.Matricula)][["Código" if "Código" in k else "Cód.", "Tipo de Conteo", "Técnico", "Creado", "Estado"]]

# --- Excel ---
COLORES = {"PENDIENTE": "F8CBAD", "Falta 2º recuento - VENCIDO": "F8CBAD", "Falta 2º": "FFE699", "Sin recuento - sin": "D9D9D9",
           "Aún no": "DDEBF7", "OK": "C6EFCE"}
HDR = PatternFill("solid", fgColor="1F3864")
wb = Workbook()


def escribir(ws, data, titulo=None):
    r0 = 1
    if titulo:
        ws.cell(1, 1, titulo).font = Font(bold=True, size=13)
        ws.cell(2, 1, f"Fecha del informe: {HOY:%d/%m/%Y} · Regla: 2 recuentos (Auditoría o Periódico) al año. "
                      f"Sin recuento y Inicio TOA antes del {LIMITE_TOA:%d/%m/%Y} = PENDIENTE.").font = Font(italic=True, size=9)
        r0 = 4
    cols = list(data.columns)
    for j, col in enumerate(cols, 1):
        cell = ws.cell(r0, j, col)
        cell.font, cell.fill = Font(bold=True, color="FFFFFF"), HDR
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for i, row in enumerate(data.itertuples(index=False), r0 + 1):
        for j, v in enumerate(row, 1):
            if isinstance(v, float) and pd.isna(v):
                v = None
            cell = ws.cell(i, j, v)
            if isinstance(v, (datetime, pd.Timestamp)):
                cell.number_format = "DD/MM/YYYY"
        if "Estado" in cols:
            est = row[cols.index("Estado")]
            for key, col in COLORES.items():
                if est.startswith(key):
                    ws.cell(i, cols.index("Estado") + 1).fill = PatternFill("solid", fgColor=col)
                    break
    for j, col in enumerate(cols, 1):
        largo = max([len(str(col))] + [len(str(v)) for v in data.iloc[:, j - 1].head(200)])
        ws.column_dimensions[get_column_letter(j)].width = min(max(10, largo + 2), 45)
    ws.freeze_panes = ws.cell(r0 + 1, 4 if "Técnico" in cols else 1)
    ws.auto_filter.ref = f"A{r0}:{get_column_letter(len(cols))}{r0 + len(data)}"


# Resumen
res = df.groupby("JDE").agg(Técnicos=("Matrícula", "count")).reset_index()
for est in ORDEN:
    res[est] = res.JDE.map(df[df.Estado == est].groupby("JDE").size()).fillna(0).astype(int)
res.insert(1, "Jefe de equipo", res.JDE.map(lambda j: nombres.get(j, j)))
ws = wb.active
ws.title = "Resumen"
escribir(ws, res, "Control de recuentos por jefe de equipo")

for jde, g in df.groupby("JDE"):
    nombre = nombres.get(jde, jde)
    ws = wb.create_sheet((nombre.split(" ")[0] + " " + nombre.split(" ")[-2] if len(nombre.split()) > 2 else nombre)[:31])
    escribir(ws, g.drop(columns="JDE"), f"Jefe de equipo: {nombre} ({jde})")

ws = wb.create_sheet("Nuevo ingreso")
_ni = nuevo_ingreso.assign(**{"Jefe de equipo": nuevo_ingreso.JDE.map(lambda j: nombres.get(j, j))})
escribir(ws, _ni[["Jefe de equipo", "Matricula", "Nombre", "Provincia", "Inicio TOA"]].fillna(""),
         "Nuevo ingreso: sin fecha de Inicio TOA en la columna O (no se envían a los jefes)")

ws = wb.create_sheet("Fuera del control")
escribir(ws, fuera, "Recuentos de personas fuera del control: filas 1-24 (PM, ADM, COORD, JDE), bajas o no activas en el cuadrante")
wb.save(salida)

# Mensaje de texto por jefe (para WhatsApp/correo)
with open(salida.replace(".xlsx", "_mensajes.txt"), "w", encoding="utf-8") as f:
    for jde, g in df.groupby("JDE"):
        nombre = nombres.get(jde, jde)
        pend = g[g.Estado.str.startswith("PENDIENTE")]
        venc = g[g.Estado == "Falta 2º recuento - VENCIDO"]
        falta = g[g.Estado == "Falta 2º recuento"]
        sin = g[g.Estado.str.startswith("Sin recuento - sin")]
        f.write(f"===== {nombre} ({jde}) =====\n")
        f.write(f"Hola {nombre.split()[0]}, te paso el control de recuentos de tu equipo a {HOY:%d/%m/%Y}. "
                f"Cada técnico debe tener 2 recuentos al año (auditoría o periódico).\n\n")
        if len(pend):
            f.write(f"🔴 SIN NINGÚN RECUENTO y más de 2 meses en TOA ({len(pend)}) – hay que hacerlo ya:\n")
            for _, t in pend.iterrows():
                f.write(f"  - {t['Matrícula']} {t['Técnico']} (Inicio TOA {t['Inicio TOA']:%d/%m/%Y})\n")
            f.write("\n")
        if len(venc):
            f.write(f"🔴 Tienen 1 recuento y ya pasaron 6 meses ({len(venc)}) – hacer el 2º ya:\n")
            for _, t in venc.iterrows():
                f.write(f"  - {t['Matrícula']} {t['Técnico']} – último {t['Último recuento']:%d/%m/%Y}, "
                        f"vencía el {t['Próximo recuento antes de']:%d/%m/%Y}\n")
            f.write("\n")
        if len(falta):
            f.write(f"🟡 Tienen 1 recuento, falta el 2º ({len(falta)}):\n")
            for _, t in falta.iterrows():
                f.write(f"  - {t['Matrícula']} {t['Técnico']} – último {t['Último recuento']:%d/%m/%Y}, "
                        f"siguiente antes del {t['Próximo recuento antes de']:%d/%m/%Y}\n")
            f.write("\n")
        if len(sin):
            f.write(f"⚪ Sin recuento y sin fecha de Inicio TOA en el cuadrante ({len(sin)}) – revisar:\n")
            for _, t in sin.iterrows():
                f.write(f"  - {t['Matrícula']} {t['Técnico']}\n")
            f.write("\n")
        ok = (g.Estado.str.startswith("OK")).sum()
        f.write(f"🟢 Al día (2 o más recuentos): {ok}\n\n")

print(df.Estado.value_counts().to_string())
print(res.to_string())
print("fuera de cuadrante:", len(fuera))

# --- Correos por jefe: Excel individual + cuerpo HTML (lo usa enviar_borradores.py) ---
import base64
import html as _html
import json
import os

dir_correos = os.path.join(os.path.dirname(os.path.abspath(salida)), f"correos_{HOY:%Y-%m-%d}")
os.makedirs(dir_correos, exist_ok=True)
SECCIONES = [  # (título, color fondo, color borde, estados, columnas [(cabecera, función)])
    ("Recuento urgente", "#FDECEA", "#C0392B",
     ["PENDIENTE - sin recuento y +2 meses en TOA", "Falta 2º recuento - VENCIDO"],
     [("Inicio TOA", lambda t: _fmt(t["Inicio TOA"])),
      ("Motivo", lambda t: "Sin ningún recuento" if t["Total recuentos"] == 0
       else f"Último {_fmt(t['Último recuento'])} · venció {_fmt(t['Próximo recuento antes de'])}")]),
    ("Falta el 2º recuento", "#FFF8E1", "#E0A800", ["Falta 2º recuento"],
     [("Último recuento", lambda t: _fmt(t["Último recuento"])),
      ("Hacer antes del", lambda t: _fmt(t["Próximo recuento antes de"]))]),
    ("Aún no les toca (menos de 2 meses en TOA)", "#EAF2FB", "#2E75B6", ["Aún no toca (menos de 2 meses en TOA)"],
     [("Inicio TOA", lambda t: _fmt(t["Inicio TOA"])),
      ("Primer recuento antes del", lambda t: _fmt(t["Próximo recuento antes de"]))]),
    ("Al día (2 o más recuentos)", "#EAF6EC", "#2E8B57", ["OK - 2 o más recuentos"],
     [("Recuentos", lambda t: " · ".join(x for x in [
         f"Aud. {t['Fechas Auditoría']}" if t["Fechas Auditoría"] else "",
         f"Per. {t['Fechas Periódico']}" if t["Fechas Periódico"] else ""] if x))]),
]


def _fmt(v):
    if v is pd.NaT:
        return ""
    if isinstance(v, (datetime, pd.Timestamp)):
        return v.strftime("%d/%m/%Y")
    return "" if v is None or (isinstance(v, float) and pd.isna(v)) else str(v)


def _tabla(g, titulo, fondo, borde, cols):
    e = _html.escape
    th = 'style="text-align:left;padding:6px 8px;font-size:12px;color:#555;border-bottom:1px solid #ddd"'
    td = 'style="padding:6px 8px;font-size:13px;border-bottom:1px solid #eee;vertical-align:top"'
    cab = f"<th {th}>Técnico</th>" + "".join(f"<th {th}>{e(c)}</th>" for c, _ in cols)
    filas = "".join(
        f"<tr><td {td}><b>{e(t['Técnico'])}</b><br><span style=\"color:#888;font-size:11px\">{e(t['Matrícula'])}</span></td>"
        + "".join(f"<td {td}>{e(f(t))}</td>" for _, f in cols) + "</tr>"
        for _, t in g.iterrows())
    return (f'<table role="presentation" width="100%" cellpadding="0" cellspacing="0" '
            f'style="border-collapse:collapse;margin:18px 0 4px;background:{fondo};border-left:4px solid {borde}">'
            f'<tr><td style="padding:8px 10px;font-size:14px;font-weight:bold;color:{borde}">{e(titulo)} ({len(g)})</td></tr></table>'
            f'<table role="presentation" width="100%" cellpadding="0" cellspacing="0" style="border-collapse:collapse">'
            f"<tr>{cab}</tr>{filas}</table>")


manifiesto = []
for jde, g in df.groupby("JDE"):
    nombre = nombres.get(jde, jde)
    archivo = f"Recuentos_{nombre.replace(' ', '_')}_{HOY:%Y-%m-%d}.xlsx"
    wbj = Workbook()
    escribir(wbj.active, g.drop(columns="JDE"), f"Jefe de equipo: {nombre} ({jde})")
    wbj.active.title = "Mi equipo"
    ruta = os.path.join(dir_correos, archivo)
    wbj.save(ruta)

    bloques, cajas = [], []
    for titulo, fondo, borde, estados, cols in SECCIONES:
        sub = g[g.Estado.isin(estados)]
        corto = titulo.split(" (")[0].split(":")[0]
        cajas.append(f'<td align="center" style="padding:8px 4px;background:{fondo};border-top:3px solid {borde}">'
                     f'<div style="font-size:22px;font-weight:bold;color:{borde}">{len(sub)}</div>'
                     f'<div style="font-size:11px;color:#555">{_html.escape(corto)}</div></td>')
        if len(sub):
            bloques.append(_tabla(sub, titulo, fondo, borde, cols))
    urgentes = int(g.Estado.isin(SECCIONES[0][3]).sum())
    cuerpo = f"""<div style="font-family:Arial,Helvetica,sans-serif;color:#222;max-width:680px">
<p style="font-size:14px">{"Buenos días," if jde in SALUDO_NEUTRO else "Hola " + _html.escape(nombre.split()[0]) + ","}</p>
<p style="font-size:14px">Te paso el control de recuentos de inventario de tu equipo a fecha <b>{HOY:%d/%m/%Y}</b>.
Cada técnico debe tener <b>2 recuentos al año</b> (auditoría o periódico). Si no tiene ninguno y lleva más de
2 meses desde su Inicio TOA, hay que hacérselo ya.</p>
<table role="presentation" width="100%" cellpadding="0" cellspacing="4" style="margin:10px 0"><tr>{''.join(cajas)}</tr></table>
{''.join(bloques)}
<p style="font-size:14px;margin-top:20px">Te adjunto el Excel con el detalle completo. Por favor, organiza los recuentos
urgentes y avísame si algún dato no cuadra.</p>
<p style="font-size:14px">Un saludo,</p>
<p style="font-size:14px;margin-top:4px">Gustavo Perez<br>Aurum BC Security SL</p></div>"""
    texto = f"Control de recuentos de tu equipo a {HOY:%d/%m/%Y}. Recuento urgente: {urgentes} técnicos. Detalle en el Excel adjunto."
    manifiesto.append({
        "jde": jde, "nombre": nombre, "asunto": f"Control de recuentos de tu equipo - {HOY:%d/%m/%Y}",
        "html": cuerpo, "texto": texto, "archivo": archivo,
        "xlsx_b64": base64.b64encode(open(ruta, "rb").read()).decode(),
    })
    # Vista previa local (documento completo con UTF-8 para que se vean bien las tildes)
    with open(os.path.join(dir_correos, archivo.replace(".xlsx", ".html")), "w", encoding="utf-8") as fh:
        fh.write('<!doctype html><html lang="es"><head><meta charset="utf-8">'
                 '<meta name="viewport" content="width=device-width,initial-scale=1">'
                 f'<title>Correo {_html.escape(nombre)}</title></head>'
                 f'<body style="margin:0;padding:16px;background:#fff">{cuerpo}</body></html>')

with open(os.path.join(dir_correos, "correos.json"), "w", encoding="utf-8") as fh:
    json.dump(manifiesto, fh, ensure_ascii=False)
print("correos preparados en:", dir_correos, "->", len(manifiesto))
