"""
PROTOTIPO: revision de la codificacion de los mantenimientos hechos el dia anterior.

  1. Descarga la ruta del dia (CSV de grupos de la consola) y lee la ficha cerrada
     de cada mantenimiento: observaciones del tecnico y lineas de codificacion.
  2. Guarda las fichas en COD_REVISION_DIR/fichas_<fecha>_<grupo>.json.
  3. (cod_reglas.py) cruza lo que dice el tecnico con lo que codifico.
  4. Un correo corto por jefe de equipo (una linea por parte mal codificado) con el detalle
     en Excel, por el Apps Script (type=recuentos-borradores): mismos destinatarios y copias
     que conteos. Registro en COD_REVISION_DIR/enviados.json para no repetir un dia.

Uso:
  python3 cod_revision.py                                  -> dia laborable anterior, toda la empresa, enviar
  python3 cod_revision.py --fecha 2026-10-26 --grupo Tenerife --modo prueba
  python3 cod_revision.py --fecha 2026-10-26 --solo-enviar  -> usa las fichas ya leidas
"""

import argparse
import json
import os
from datetime import date, timedelta
from pathlib import Path

import consola_sso
import ofs_rm_prueba
import ruta_sugerencias
from comun_ofs import LAUNCH_ARGS
from mac_tarde import SESION_MAC, _entrar

CARPETA = Path(os.environ.get("COD_REVISION_DIR", "/Users/gustavoperez/Correo Recuentos/cod_revision"))


def dia_anterior(hoy=None):
    hoy = hoy or date.today()
    d = hoy - timedelta(days=1)
    while d.weekday() >= 5:   # lunes revisa el viernes
        d -= timedelta(days=1)
    return d


def leer(fecha_iso, grupo=""):
    from playwright.sync_api import sync_playwright
    correo, clave = consola_sso.leer_credenciales_sso()
    errores, capturas, datos = [], [], []
    with sync_playwright() as p:
        browser = p.chromium.launch(headless=True, args=LAUNCH_ARGS)
        ctx, page, _ = _entrar(browser, correo, clave)
        try:
            csvs = [c for c in consola_sso.descargar_csv_grupos(page, fecha_iso) if grupo.lower() in c[0].lower()]
            filas = ruta_sugerencias._filas_ruta(csvs, fecha_iso)
            print(f"[cod] {len(filas)} mantenimientos el {fecha_iso} ({grupo or 'todos'})", flush=True)
            for i, f in enumerate(filas, 1):
                ficha = ofs_rm_prueba._leer(page, f["numero"], capturas, errores)
                datos.append(dict(f, ficha=ficha if isinstance(ficha, dict) else None))
                if i % 10 == 0:
                    print(f"[cod] {i}/{len(filas)}", flush=True)
            ctx.storage_state(path=str(SESION_MAC))
        finally:
            browser.close()
    CARPETA.mkdir(parents=True, exist_ok=True)
    salida = CARPETA / f"fichas_{fecha_iso}_{grupo or 'todos'}.json"
    salida.write_text(json.dumps(datos, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"[cod] {sum(1 for d in datos if d['ficha'])}/{len(datos)} fichas leidas ({len(errores)} errores) -> {salida}")
    return datos


REGISTRO = CARPETA / "enviados.json"


def _registro():
    try:
        return json.loads(REGISTRO.read_text(encoding="utf-8"))
    except Exception:
        return {}


def _excel(filas, destino):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    wb = Workbook()
    ws = wb.active
    ws.title = "Partes"
    ws.append(["Técnico", "Matrícula", "Mantenimiento", "Tipo", "Qué corregir", "Qué escribió el técnico", "Qué codificó"])
    for f in filas:
        ws.append(f)
    for c in ws[1]:
        c.font = Font(bold=True)
    for col, ancho in zip("ABCDEFG", (26, 11, 13, 22, 60, 70, 50)):
        ws.column_dimensions[col].width = ancho
    for fila in ws.iter_rows(min_row=2):
        for c in fila:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(destino)
    return destino


def revisar_y_mandar(fecha_iso, datos, modo="enviar"):
    import base64
    import re
    from collections import defaultdict
    import cod_reglas as cr
    from comun_ofs import enviar_aviso
    from ov_conteos import _jerarquia, _mandar

    if modo == "enviar" and fecha_iso in _registro():
        print(f"[cod] el {fecha_iso} ya se envio ({_registro()[fecha_iso]})")
        return
    jer = _jerarquia()
    por_je = defaultdict(lambda: defaultdict(list))
    fin_je = defaultdict(lambda: defaultdict(int))   # JE -> matricula -> partes FINL
    total = finl = 0
    datos = list({d["numero"]: d for d in datos}.values())   # un parte puede salir en dos grupos de la ruta
    for d in datos:
        f = d.get("ficha") or {}
        if (f.get("estado") or "").upper() != "FINL":
            continue
        finl += 1
        mat = (d.get("tecnico") or f.get("tecnico") or "").upper()
        p = jer.get(mat) or jer.get(re.sub(r"^[A-Z]{2}(?=\d{6}$)", "", mat)) or {}
        je = (p.get("JE") or "").strip().upper() or "SIN JE"
        if je == "302154":
            je = "GD5381"   # los de Nicola Morsa: a Gustavo con Nicola en CCO (CORREO_FORZADO / CCO_JEFE del Apps Script)
        fin_je[je][mat] += 1
        h = cr.revisar(f)
        if not h:
            continue
        total += 1
        por_je[je][(p.get("Nombre") or mat, mat)].append((d["numero"], f, h))
    fecha_txt = "/".join(reversed(fecha_iso.split("-")))
    jefes, resumen = [], []
    for je, tecs in por_je.items():
        if je == "SIN JE":
            resumen.append(f"Sin jefe de equipo en el cuadrante: {', '.join(m for _, m in tecs)}")
            continue
        html = cr.correo_html_corto(fecha_txt, [(n, m, [(num, h) for num, _, h in ms]) for (n, m), ms in tecs.items()],
                                    fin_je[je])
        filas = [[n, m, num, f.get("tipo"), " · ".join(dict.fromkeys(cr.corto(x["texto"]) for x in h)),
                  cr._observaciones(f), cr._cod_texto(f)] for (n, m), ms in tecs.items() for num, f, h in ms]
        archivo = f"Codificacion {fecha_iso} JE {je}.xlsx"
        _excel(filas, CARPETA / archivo)
        n_p = sum(len(ms) for ms in tecs.values())
        jefes.append({"codigo": je, "jde": je, "asunto": f"Partes del {fecha_txt} con la codificación mal ({n_p})",
                      "html": html, "texto": f"Partes del {fecha_txt} con la codificación mal: {n_p}.",
                      "archivo": archivo, "xlsx_b64": base64.b64encode((CARPETA / archivo).read_bytes()).decode()})
        n_f = sum(fin_je[je].values())
        resumen.append(f"JE {je}: {n_p} de {n_f} partes mal ({cr._pct(n_p, n_f)}), {len(tecs)} técnicos")
    errores = []
    if jefes:
        res = _mandar(jefes, modo)
        ok = sum(1 for r in res.get("resultado", []) if r.get("ok"))
        errores = [f"JE {r.get('jde')}: {r.get('error')}" for r in res.get("resultado", []) if not r.get("ok")]
        resumen.insert(0, f"{fecha_txt}: {finl} partes finalizados, {total} con la codificación mal ({cr._pct(total, finl)}). "
                          f"Modo {res.get('modo')}: {ok}/{len(jefes)} correos ok.")
    if modo == "enviar" and jefes:
        from datetime import datetime
        reg = _registro()
        reg[fecha_iso] = datetime.now().strftime("%Y-%m-%d %H:%M")
        REGISTRO.write_text(json.dumps(reg, indent=1), encoding="utf-8")
    if modo != "prueba":
        enviar_aviso(("⚠️" if errores else "✅") + f" Codificación de partes del {fecha_txt}",
                     "\n".join(resumen + ([""] + ["Errores:"] + errores if errores else [])))
    print("\n".join(resumen + errores))


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--fecha")
    ap.add_argument("--grupo", default="")
    ap.add_argument("--modo", choices=["enviar", "prueba", "borradores"], default="enviar")
    ap.add_argument("--solo-leer", action="store_true")
    ap.add_argument("--solo-enviar", action="store_true")
    a = ap.parse_args()
    fecha = a.fecha or dia_anterior().isoformat()
    if a.modo == "enviar" and not a.solo_leer and fecha in _registro():
        print(f"[cod] los partes del {fecha} ya se enviaron ({_registro()[fecha]}): nada que hacer")
        raise SystemExit(0)
    if a.solo_enviar:
        datos = json.loads((CARPETA / f"fichas_{fecha}_{a.grupo or 'todos'}.json").read_text(encoding="utf-8"))
    else:
        datos = leer(fecha, a.grupo)
    if not a.solo_leer:
        revisar_y_mandar(fecha, datos, a.modo)
